import pytest


@pytest.fixture
def sales(monkeypatch,tmp_path):
    from db import database,crud
    from tools import report_tool,office_tool
    monkeypatch.setattr(database,'_DATA_DIR',tmp_path)
    monkeypatch.setattr(database,'DB_PATH',tmp_path/'agent.db')
    monkeypatch.setattr(report_tool,'REPORTS_DIR',tmp_path/'reports')
    monkeypatch.setattr(office_tool,'OFFICE_DIR',tmp_path/'reports')
    database.init_db()
    doc=crud.insert_document('synthetic-sales.png')
    crud.insert_rows(doc,[{'desc':'合成甲公司','date':'2026-10-03','item':'A','amount':'2','price':'999','sum':'100'},
                          {'desc':'合成甲公司','date':'2026-10-03','item':'折扣','amount':'0','sum':'-25'},
                          {'desc':'合成乙公司','date':'2026-10-03','item':'B','amount':'3','sum':'240'}])
    crud.confirm_doc(doc)
    return tmp_path


def test_word_respects_keyword_and_persists_snapshot(sales):
    from tools.report_tool import generate_report
    from db import artifacts
    from docx import Document
    output=generate_report.invoke({'year':'2026','keyword':'合成甲','include_chart':False})
    record=artifacts.get(output.cards[0]['artifact_id'])
    assert record['snapshot']['summary']['total_sum']==75
    overview={r.cells[0].text:r.cells[1].text for r in Document(record['path']).tables[0].rows}
    assert overview['总销售额']=='75.00'
    assert overview['确认记录行数']=='2 行'


def test_office_formats_reuse_identical_snapshot_and_excel_is_complete(sales):
    from tools.report_tool import generate_report
    from tools.office_tool import generate_excel,generate_presentation
    from db import artifacts,crud
    from openpyxl import load_workbook
    from pptx import Presentation
    word=generate_report.invoke({'year':'2026','keyword':'合成甲','include_chart':False})
    ident=word.cards[0]['artifact_id']
    # Change live business data; explicit snapshot reuse must keep export scope frozen.
    crud.update_row(1,{'sum':'999'})
    excel=generate_excel.invoke({'snapshot_id':ident})
    ppt=generate_presentation.invoke({'snapshot_id':ident})
    records=[artifacts.get(x.cards[0]['artifact_id']) for x in (word,excel,ppt)]
    assert all(r['snapshot']==records[0]['snapshot'] for r in records)
    workbook=load_workbook(records[1]['path'],data_only=True)
    assert workbook['销售明细'].max_row==3
    assert sum(row[7] for row in workbook['销售明细'].iter_rows(min_row=2,values_only=True))==75
    workbook.close()
    deck=Presentation(records[2]['path'])
    assert len(deck.slides)>=5
    assert any('75.00' in shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)


def test_empty_scope_produces_no_office_file(sales):
    from tools.office_tool import generate_excel,generate_presentation
    for tool in [generate_excel,generate_presentation]:
        output=tool.invoke({'year':'1999'})
        assert '没有符合条件' in output
        assert not getattr(output,'cards',[])
