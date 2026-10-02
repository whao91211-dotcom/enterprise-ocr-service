import pytest


@pytest.fixture
def artifacts(monkeypatch, tmp_path):
    from db import database
    from tools import plot_tool, report_tool
    from evals.real_text_baseline import _seed_sales
    monkeypatch.setattr(database, '_DATA_DIR', tmp_path)
    monkeypatch.setattr(database, 'DB_PATH', tmp_path/'agent.db')
    monkeypatch.setattr(plot_tool, 'CHARTS_DIR', tmp_path/'charts')
    monkeypatch.setattr(report_tool, 'REPORTS_DIR', tmp_path/'reports')
    _seed_sales()
    return plot_tool, report_tool, tmp_path


def test_report_year_applies_to_document_count_and_totals(artifacts):
    from docx import Document
    _, report_tool, tmp = artifacts
    report_tool.generate_report.invoke({'year': '2024', 'include_chart': False})
    document = Document(next((tmp/'reports').glob('销售报告_2024*.docx')))
    overview = {row.cells[0].text: row.cells[1].text for row in document.tables[0].rows}
    assert overview['涉及单据'] == '2 张'
    assert overview['确认记录行数'] == '2 行'
    assert overview['总销售额'] == '300.00'


@pytest.mark.parametrize('arguments', [{'group_by': 'unsupported'}, {'chart_type': 'line'}])
def test_chart_rejects_unsupported_options_instead_of_silent_default(artifacts, arguments):
    plot_tool, _, tmp = artifacts
    result = plot_tool.plot_chart.invoke(arguments)
    assert result.startswith(('错误:', 'group_by 仅支持:'))
    assert not list((tmp/'charts').glob('*.png'))


def test_zero_total_pie_has_clear_error_and_no_artifact(artifacts):
    from db import crud
    plot_tool, _, tmp = artifacts
    doc = crud.insert_document('zero.png')
    crud.insert_rows(doc, [{'desc': '零公司', 'date': '2026-01-01', 'amount': '1', 'sum': '0'}])
    crud.confirm_doc(doc)
    result = plot_tool.plot_chart.invoke({'chart_type': 'pie', 'filter_query': '零公司', 'year': '2026'})
    assert result.startswith('错误:')
    assert '饼图' in result
    assert not list((tmp/'charts').glob('*.png'))


def test_report_discloses_chart_failure_in_artifact_and_tool_response(artifacts, monkeypatch):
    from docx import Document
    plot_tool, report_tool, tmp = artifacts

    def broken(*args, **kwargs):
        raise RuntimeError('synthetic chart failure')

    monkeypatch.setattr(plot_tool, '_collect_data', broken)
    result = report_tool.generate_report.invoke({'year': '2024', 'include_chart': True})
    document = Document(next((tmp/'reports').glob('销售报告_2024*.docx')))
    assert '图表未生成' in result
    assert any('图表未生成' in p.text for p in document.paragraphs)
    assert len(document.inline_shapes) == 0


def test_repeated_reports_preserve_previous_snapshot(artifacts):
    from db import crud
    from docx import Document
    _, report_tool, tmp = artifacts
    report_tool.generate_report.invoke({'year': '2024', 'include_chart': False})
    first = next((tmp/'reports').glob('*.docx'))
    doc = crud.insert_document('new-sale.png')
    crud.insert_rows(doc, [{'desc': '甲公司', 'date': '2024-06-01', 'amount': '1', 'sum': '50'}])
    crud.confirm_doc(doc)
    report_tool.generate_report.invoke({'year': '2024', 'include_chart': False})
    assert len(list((tmp/'reports').glob('*.docx'))) == 2
    overview = {r.cells[0].text: r.cells[1].text for r in Document(first).tables[0].rows}
    assert overview['总销售额'] == '300.00'


def test_chart_variants_preserve_both_outputs(artifacts):
    plot_tool, _, tmp = artifacts
    plot_tool.plot_chart.invoke({'year': '2024', 'chart_type': 'bar'})
    plot_tool.plot_chart.invoke({'year': '2024', 'chart_type': 'pie'})
    assert len(list((tmp/'charts').glob('*.png'))) == 2


def test_chart_save_failure_closes_figure(artifacts, monkeypatch):
    from matplotlib.figure import Figure
    plot_tool, _, _ = artifacts
    before = plot_tool.plt.get_fignums()

    def fail(*args, **kwargs):
        raise OSError('synthetic disk failure')

    monkeypatch.setattr(Figure, 'savefig', fail)
    with pytest.raises(OSError):
        plot_tool.plot_chart.invoke({})
    assert plot_tool.plt.get_fignums() == before


def test_report_picture_failure_removes_intermediate_png(artifacts, monkeypatch):
    from docx.document import Document
    _, report_tool, tmp = artifacts

    def fail(*args, **kwargs):
        raise ValueError('synthetic picture embedding failure')

    monkeypatch.setattr(Document, 'add_picture', fail)
    assert '图表未生成' in report_tool.generate_report.invoke({'include_chart': True})
    assert not list((tmp/'reports').glob('*.png'))
