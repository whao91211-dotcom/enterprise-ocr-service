"""Reproducible local product/Office probe; only generated synthetic business data."""
import argparse
import json
import os
import sys
import time
from pathlib import Path

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--root',type=Path,required=True)
    parser.add_argument('--office-preview',action='store_true')
    args=parser.parse_args()
    root=args.root.resolve()
    root.mkdir(parents=True,exist_ok=True)
    os.environ['AGENT_DATA_DIR']=str(root)
    from db import database,crud,review,artifacts,chat_memory
    from tools import report_tool,office_tool
    from services.office_preview import convert_artifact
    from openpyxl import load_workbook
    from pptx import Presentation
    from docx import Document
    database.init_db()
    if crud.all_docs_with_stats():
        raise SystemExit('Use a fresh output directory; preserve prior run results.')
    report_tool.REPORTS_DIR=office_tool.OFFICE_DIR=root/'files'
    doc=crud.insert_document('synthetic-product.png')
    rows=[{'desc':'合成销售公司（办公预览测试）','from':'合成供货方','date':'2026-10-03',
           'item':f'办公商品{i:02d}','amount':'2','price':'50','tax':'10%','sum':'100'} for i in range(18)]
    rows.append({'desc':'其他公司','date':'2026-10-03','item':'排除商品','amount':'9','sum':'900'})
    crud.insert_rows(doc,rows)
    sid,_=chat_memory.resolve_session(None)
    current=review.get_review(doc)
    start=time.perf_counter()
    edited=review.save_edits(doc,current['version'],[{'id':current['rows'][0]['id'],'fields':{'sum':'150'}}])
    edit_ms=round((time.perf_counter()-start)*1000,2)
    pending_after_edit=all(r['status']=='pending' for r in edited['rows'])
    review.confirm(doc,edited['version'])
    outputs=[]
    for tool,params in [(report_tool.generate_report,{'year':'2026','keyword':'合成销售','group_by':'desc'}),
                        (office_tool.generate_excel,{}),(office_tool.generate_presentation,{'group_by':'desc'})]:
        if outputs:
            params['snapshot_id']=outputs[0].cards[0]['artifact_id']
        outputs.append(tool.invoke(params))
    cards=[o.cards[0] for o in outputs]
    chat_memory.save_turn(sid,'生成销售办公材料','三种文件已生成，请预览核对。',assistant_meta={'cards':cards+[{'type':'ocr_review','doc_id':doc}]})
    records=[artifacts.get(c['artifact_id']) for c in cards]
    wb=load_workbook(records[1]['path'],data_only=True)
    excel_rows=wb['销售明细'].max_row-1
    excel_sum=sum(r[7] for r in wb['销售明细'].iter_rows(min_row=2,values_only=True))
    wb.close()
    word=Document(records[0]['path'])
    overview={r.cells[0].text:r.cells[1].text for r in word.tables[0].rows}
    deck=Presentation(records[2]['path'])
    checks={'edit_keeps_pending':pending_after_edit,'same_snapshot':all(r['snapshot']==records[0]['snapshot'] for r in records),
            'filtered_rows':excel_rows==18,'excel_total':excel_sum==1850,'word_total':overview['总销售额']=='1,850.00' or overview['总销售额']=='1850.00',
            'ppt_total':any('1,850.00' in s.text or '1850.00' in s.text for slide in deck.slides for s in slide.shapes if s.has_text_frame),
            'history_cards':len(chat_memory.message_page(sid)['messages'][-1]['meta']['cards'])==4}
    previews=[]
    if args.office_preview:
        import pypdfium2
        for record in (records[0],records[2]):
            start=time.perf_counter()
            convert_artifact(record['id'])
            fresh=artifacts.get(record['id'])
            item={'kind':record['kind'],'status':fresh['preview_status'],'elapsed_ms':round((time.perf_counter()-start)*1000)}
            if fresh['preview_status']=='ready':
                pdf=pypdfium2.PdfDocument(fresh['preview_path'])
                item['pages']=len(pdf)
                images=[]
                for i in range(len(pdf)):
                    page=pdf[i]
                    bitmap=page.render(scale=1.2)
                    image_path=root/f"{record['kind']}-page-{i+1}.png"
                    bitmap.to_pil().save(image_path)
                    images.append(str(image_path))
                    bitmap.close()
                    page.close()
                pdf.close()
                item['images']=images
            previews.append(item)
            checks[record['kind']+'_preview']=item['status']=='ready'
    result={'mode':'offline synthetic local Office probe','session_id':sid,'checks':checks,'direct_edit_ms':edit_ms,
            'direct_edit_llm_calls':0,'files':[{k:r[k] for k in ('id','kind','name')} for r in records],'previews':previews,
            'limits':['One local run, no comparable prior UI timing; not a latency improvement claim',
                      'Synthetic sales only; not general Office compatibility or overall success rate']}
    (root/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)
    if not all(checks.values()):
        raise SystemExit(1)


if __name__=='__main__':
    main()
