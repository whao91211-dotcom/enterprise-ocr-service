"""Prepare reviewable OCR labels and Agent task candidates; never calls services."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

FIELDS = ['desc', 'date', 'from', 'item', 'amount', 'price', 'tax', 'sum']
NUMERIC = {'amount', 'price', 'tax', 'sum'}


def number(value):
    text = str(value).strip().replace('¥', '').replace('￥', '').replace(',', '')
    percent = text.endswith('%')
    result = Decimal(text.rstrip('%'))
    if not result.is_finite():
        raise InvalidOperation
    return result / 100 if percent else result


def normalize_row(raw, excel_row):
    values = {}
    for field, value in zip(FIELDS, raw[1:9]):
        if value is None or value == '':
            values[field] = None
        elif field in NUMERIC:
            try:
                values[field] = format(number(value).normalize(), 'f')
            except InvalidOperation:
                values[field] = str(value)
        elif field == 'date':
            text = str(value)
            match = re.fullmatch(r'(\d{4})年(\d{1,2})月(\d{1,2})日', text)
            if match:
                text = '{}-{:02d}-{:02d}'.format(int(match[1]), int(match[2]), int(match[3]))
            try:
                values[field] = datetime.fromisoformat(text).date().isoformat()
            except ValueError:
                values[field] = text
        else:
            values[field] = str(value)
    return {'excel_row': excel_row,
            'image_name': str(raw[0] or '').replace('\\', '/').split('/')[-1],
            'raw': [str(v) if isinstance(v, datetime) else v for v in raw[:9]],
            'values': values}


def inspect_rows(rows):
    issues, seen = [], {}
    for row in rows:
        def issue(code, field='', detail=''):
            issues.append({'image_name': row['image_name'],
                           'excel_row': row['excel_row'], 'code': code,
                           'field': field, 'detail': detail, 'review_decision': ''})
        v = row['values']
        for field, value in v.items():
            if value is None:
                issue('blank_field', field, '核对原图是否为空；空值不等于零')
            elif field in NUMERIC:
                try:
                    number(value)
                except InvalidOperation:
                    issue('invalid_number', field, str(value))
            elif field == 'date':
                try:
                    datetime.strptime(value, '%Y-%m-%d')
                except ValueError:
                    issue('invalid_date', field, str(value))
        try:
            if all(v[k] is not None for k in ('amount', 'price', 'sum')):
                if abs(number(v['amount']) * number(v['price']) - number(v['sum'])) > Decimal('.01'):
                    issue('amount_requires_image_check', 'sum', '可能有折扣；不自动改金额')
        except InvalidOperation:
            pass
        key = (row['image_name'], tuple(v.get(k) for k in FIELDS))
        if key in seen:
            issue('repeated_row_preserved', '', f"与 Excel 行 {seen[key]} 相同，保留并核对原图")
        else:
            seen[key] = row['excel_row']
    return issues


def dump(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')


def prepare(source: Path, output: Path):
    import openpyxl  # read-only extraction; optional dataset preparation dependency
    from PIL import Image

    workbook_path = source / 'data.xlsx'
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=False)
    cached_workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook['Sheet1']
    expected = ['图像名', '顾客公司', '发注日', '源公司', '项目', '数量', '单价', '税率', '金额']
    if [sheet.cell(1, i).value for i in range(1, 10)] != expected:
        raise ValueError('Unexpected workbook header')
    rows, extra_issues = [], []
    paired = zip(sheet.iter_rows(min_row=2, values_only=True),
                 cached_workbook['Sheet1'].iter_rows(min_row=2, values_only=True))
    for i, (raw, cached) in enumerate(paired, 2):
        if not any(v is not None for v in raw):
            continue
        effective = [cached[c] if isinstance(v, str) and v.startswith('=') else v
                     for c, v in enumerate(raw)]
        row = normalize_row(effective, i)
        row['raw'] = [str(v) if isinstance(v, datetime) else v for v in raw[:9]]
        rows.append(row)
        for col, value in enumerate(raw, 1):
            if (isinstance(value, str) and value.startswith('=')) or (col > 9 and value is not None):
                extra_issues.append({'image_name': rows[-1]['image_name'], 'excel_row': i,
                                     'code': 'formula_derived_label' if col <= 9 else 'extra_column', 'field': str(col),
                                     'detail': f'原值: {value}; cached={cached[col-1]}', 'review_decision': ''})
    workbook.close()
    cached_workbook.close()
    issues = inspect_rows(rows) + extra_issues
    images = {p.name: p for p in (source / 'sales').iterdir() if p.is_file()}
    grouped = {}
    for row in rows:
        grouped.setdefault(row['image_name'], []).append(row)
    samples, hashes = [], {}
    for name, records in sorted(grouped.items()):
        path = images.get(name)
        meta = {'sample_id': Path(name).stem, 'image_name': name,
                'image_relative_path': f'sales/{name}', 'review_status': 'pending',
                'rows': records, 'image_valid': False}
        if path:
            meta['image_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
            hashes.setdefault(meta['image_sha256'], []).append(name)
            try:
                with Image.open(path) as image:
                    meta['image_size'] = list(image.size)
                    image.verify()
                meta['image_valid'] = True
            except Exception as exc:
                issues.append({'image_name': name, 'excel_row': '', 'code': 'unreadable_image',
                               'field': '', 'detail': str(exc), 'review_decision': ''})
        else:
            issues.append({'image_name': name, 'excel_row': '', 'code': 'missing_image',
                           'field': '', 'detail': '', 'review_decision': ''})
        samples.append(meta)
    for names in hashes.values():
        if len(names) > 1:
            issues.append({'image_name': '|'.join(names), 'excel_row': '', 'code': 'identical_images',
                           'field': '', 'detail': '相同文件内容；后续划分时应放在同一分组', 'review_decision': ''})
    # Agent fixtures exclude whole documents with missing/invalid numbers or dates.
    blockers = {'blank_field', 'invalid_number', 'invalid_date', 'formula_derived_label', 'extra_column'}
    excluded = {i['image_name'] for i in issues if i['code'] in blockers}
    fixture_samples = [s for s in samples if s['image_valid'] and s['image_name'] not in excluded][:20]
    fixture_rows = [r for s in fixture_samples for r in s['rows']]
    tasks = []
    for field in ('desc', 'item'):
        keywords = sorted({r['values'][field] for r in fixture_rows})[:10]
        for keyword in keywords:
            matching = [r for r in fixture_rows if any(keyword in (r['values'][k] or '')
                        for k in ('desc', 'from', 'item')) or keyword in r['image_name']]
            tasks.append({'case_id': f'query-{len(tasks)+1:02d}', 'review_status': 'pending',
                          'question': f'在已确认数据中，关键词“{keyword}”匹配的总金额和明细行数是多少？',
                          'expected': {'keyword': keyword, 'matched_rows': len(matching),
                                       'total_sum': str(sum((number(r['values']['sum']) for r in matching), Decimal(0)))},
                          'source_excel_rows': [r['excel_row'] for r in matching],
                          'acceptance': '必须按关键词在商品/公司/文件名中匹配并统计全量；不能用20行样本求总额。'})
    no_match = '__benchmark_no_match_20260930__'
    assert not any(no_match in str(r['raw']) for r in fixture_rows)
    tasks.append({'case_id': 'query-no-match', 'review_status': 'pending',
                  'question': f'查询关键词“{no_match}”的已确认销售数据。',
                  'expected': {'keyword': no_match, 'matched_rows': 0},
                  'source_excel_rows': [], 'acceptance': '明确没有匹配数据，不编造结果。'})
    output.mkdir(parents=True, exist_ok=True)
    with (output / 'ocr_samples.jsonl').open('w', encoding='utf-8') as f:
        for sample in samples:
            f.write(json.dumps(sample, ensure_ascii=False, default=str) + '\n')
    with (output / 'review.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=['image_name', 'excel_row', 'code', 'field', 'detail', 'review_decision'])
        writer.writeheader()
        writer.writerows(issues)
    dump(output / 'agent_fixture.json', {'review_status': 'pending', 'status_for_evaluation': 'confirmed',
         'scope': 'isolated evaluation database only; never production', 'samples': fixture_samples})
    dump(output / 'agent_tasks.json', {'review_status': 'pending', 'fixture': 'agent_fixture.json', 'cases': tasks})
    with (output / 'task_review.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(['case_id', 'question', 'matched_rows', 'total_sum', 'source_excel_rows', 'review_decision'])
        for task in tasks:
            writer.writerow([task['case_id'], task['question'], task['expected']['matched_rows'],
                             task['expected'].get('total_sum', ''),
                             '|'.join(map(str, task['source_excel_rows'])), ''])
    summary = {'schema_version': 1, 'source_workbook': str(workbook_path.resolve()),
               'workbook_sha256': hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
               'training_overlap': 'user reports no overlap; not independently verified',
               'review_status': 'pending', 'label_rows': len(rows), 'labelled_images': len(samples),
               'image_files': len(images), 'valid_labelled_images': sum(s['image_valid'] for s in samples),
               'unlabelled_images': sorted(set(images)-set(grouped)),
               'issue_counts': dict(Counter(i['code'] for i in issues)),
               'issue_image_counts': {code: len({i['image_name'] for i in issues if i['code'] == code})
                                      for code in sorted({i['code'] for i in issues})},
               'blank_fields': dict(Counter(i['field'] for i in issues if i['code'] == 'blank_field')),
               'agent_fixture_images': len(fixture_samples), 'agent_fixture_rows': len(fixture_rows),
               'agent_task_candidates': len(tasks),
               'limits': ['program checks are not human label verification',
                          'discount discrepancies and repeated rows are review hints, not proven errors',
                          'no OCR or Agent execution; no quality scores',
                          'normal query tasks only; fault injection and multi-turn cases not yet implemented']}
    dump(output / 'validation.json', summary)
    guide = ['# Benchmark v1 人工核查入口', '',
             '状态：待人工核查。程序校验通过不代表标注已被人工逐图确认。', '',
             f"共 {len(samples)} 张图、{len(rows)} 行；首批 Agent 夹具 {len(fixture_samples)} 张图、{len(fixture_rows)} 行，{len(tasks)} 个查询任务。", '',
             '## 优先核查', '',
             '| 类型 | 提示数 | 涉及图片 | 处理原则 |', '| --- | ---: | ---: | --- |']
    descriptions = {'formula_derived_label': '优先：核对图片原金额；公式缓存值不是独立人工答案',
                    'blank_field': '核对图片是否确实为空；不要直接填零',
                    'amount_requires_image_check': '核对折扣与原金额；不要直接重算覆盖',
                    'repeated_row_preserved': '核对重复次数；不要自动去重'}
    for code, count in summary['issue_counts'].items():
        guide.append(f"| {code} | {count} | {summary['issue_image_counts'][code]} | {descriptions.get(code, '核对原始文件')} |")
    guide += ['', '提示可能重叠，不能相加当作错误行数。', '',
              '## 文件与核查方式', '',
              '- [逐项核查清单](review.csv)：根据 image_name 打开原图，excel_row 定位原表。review_decision 可填“确认原值 / 应修正为… / 暂不能确认”。',
              '- [任务核查清单](task_review.csv)：核查问题与标准金额；任务仅使用 agent_fixture.json 的隔离数据，不使用生产数据库。',
              '- [OCR 样本](ocr_samples.jsonl)：保留原值、规范化值、Excel 行号和图片哈希；所有样本待人工核查。',
              '- [Agent 数据夹具](agent_fixture.json) 与 [任务定义](agent_tasks.json)：未执行模型，不是任务成功率报告。',
              '- [校验统计](validation.json)。', '',
              '本版本不划分调参/最终测试集。之后按图片分组划分，不能把同图不同行分散到两组。',
              'OCR 金标准以图片实际内容为准，金额不能由数量和单价自动推导。公式缓存可能过期。',
              '日期仅做 ISO / 年月日格式统一，货币数字去符号与千分位；不翻译公司或商品名。',
              '空字段的评分口径需要人工确认后决定；重复行评分必须保留次数。',
              '异常注入、多轮记忆和完整 Agent 执行评估将另行补充，本版本只准备数据。']
    (output / 'REVIEW.md').write_text('\n'.join(guide) + '\n', encoding='utf-8')
    return summary


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=Path('data/benchmark_v1'))
    args = parser.parse_args()
    print(json.dumps(prepare(args.source, args.output), ensure_ascii=False, indent=2))
