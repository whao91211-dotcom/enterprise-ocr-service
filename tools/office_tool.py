"""Fixed, editable sales workbooks and presentations from registered snapshots."""
from pathlib import Path
from uuid import uuid4
from langchain_core.tools import tool
from agent.runtime_guard import ToolOutcome
from db import artifacts,crud
from services.sales_snapshot import capture, number

OFFICE_DIR=Path(__file__).resolve().parents[1]/'reports'
LABELS=['顾客公司','发注日','源公司','商品','数量','单价','税率','票面金额']


def _snapshot(year,keyword,snapshot_id,group_by):
    if group_by not in ('item','desc','date'):
        raise ValueError('group_by 仅支持 item、desc、date')
    return capture(year,keyword,snapshot_id)


def _result(path,kind,snapshot):
    ident=artifacts.register(path,kind,snapshot)
    label='Excel' if kind=='xlsx' else 'PPT'
    return ToolOutcome(f'{label}已生成: {path}\n统计范围：{snapshot["filters"]}；'
        f'票面金额合计 {snapshot["summary"]["total_sum"]:.2f}；snapshot_id={ident}。文件卡片提供预览与下载。',
        cards=[{'type':'artifact','artifact_id':ident}])


@tool
def generate_excel(year: str='',keyword: str='',group_by: str='item',snapshot_id: str='') -> str:
    """生成已确认销售数据的 Excel 工作簿，含完整明细、概览、三类分组和图表。

    Args:
        year: 四位年份；空为全部年份。
        keyword: 与查询工具相同的公司/商品/源公司/文件名关键词。
        group_by: 图表分组 item/desc/date。
        snapshot_id: 可复用此前办公文件返回的 snapshot_id，确保多格式数据一致；空则重新取最新数据。
    """
    try:
        snapshot=_snapshot(year,keyword,snapshot_id,group_by)
    except (ValueError,LookupError,FileNotFoundError) as exc:
        return '错误: '+str(exc)
    if not snapshot['rows']:
        return '（没有符合条件的已确认数据，无法生成 Excel）'
    import xlsxwriter
    OFFICE_DIR.mkdir(parents=True,exist_ok=True)
    path=OFFICE_DIR/f'销售明细_{uuid4().hex}.xlsx'
    with xlsxwriter.Workbook(str(path),{'strings_to_formulas':False,'strings_to_urls':False}) as workbook:
        header=workbook.add_format({'bold':True,'bg_color':'#24443C','font_color':'white','border':1,'text_wrap':True})
        money=workbook.add_format({'num_format':'#,##0.00;[Red]-#,##0.00'})
        integer=workbook.add_format({'num_format':'0.##'})
        summary=workbook.add_worksheet('统计概览');summary.set_column('A:A',24);summary.set_column('B:B',40)
        summary.write_row(0,0,['统计项目','数值或范围'],header)
        scope=snapshot['filters']
        for i,(key,value) in enumerate([('年份',scope['year'] or '全部'),('关键词',scope['keyword'] or '全部'),
            ('涉及单据',snapshot['summary']['total_docs']),('确认记录行数',snapshot['summary']['total_rows']),
            ('总销售额',snapshot['summary']['total_sum']),('总销售数量',snapshot['summary']['total_amount']),
            ('快照时间',snapshot['created_at']),('口径','金额以票面金额为准；未确认数据不计入')],1):
            summary.write(i,0,key);summary.write(i,1,value,money if key=='总销售额' else None)
        ws=workbook.add_worksheet('销售明细');ws.freeze_panes(1,0);ws.set_column(0,3,23);ws.set_column(4,7,16);ws.set_column(8,9,26)
        ws.write_row(0,0,LABELS+['源文件','行 ID'],header)
        for i,row in enumerate(snapshot['rows'],1):
            for j,key in enumerate(crud.USER_FIELDS):
                value=row.get(key) or ''
                if j in (4,5,7) and value:
                    ws.write_number(i,j,float(number(value)),money if j in (5,7) else integer)
                else:
                    ws.write_string(i,j,str(value))
            ws.write_string(i,8,row['file_name']);ws.write_number(i,9,row['id'])
        ws.autofilter(0,0,len(snapshot['rows']),9)
        sheet_names={'item':'按商品汇总','desc':'按公司汇总','date':'按日期汇总'}
        for key,groups in snapshot['groups'].items():
            tab=workbook.add_worksheet(sheet_names[key]);tab.set_column(0,0,30);tab.set_column(1,3,18);tab.freeze_panes(1,0)
            tab.write_row(0,0,['分组','数量合计','票面金额合计','明细行数'],header)
            for i,g in enumerate(groups,1):
                tab.write_string(i,0,g['group']);tab.write_number(i,1,g['amount'],integer);tab.write_number(i,2,g['total'],money);tab.write_number(i,3,g['rows'])
            tab.autofilter(0,0,len(groups),3)
        selected=snapshot['groups'][group_by]
        chart=workbook.add_chart({'type':'column'})
        chart.add_series({'name':'票面金额','categories':[sheet_names[group_by],1,0,min(15,len(selected)),0],
                          'values':[sheet_names[group_by],1,2,min(15,len(selected)),2],'fill':{'color':'#496C60'}})
        chart.set_title({'name':'销售金额分组（前15组）'});chart.set_legend({'none':True})
        summary.insert_chart('D2',chart)
        summary.write(10,0,'图表最多展示前15组，各汇总表与明细包含全部匹配记录。')
    return _result(path,'xlsx',snapshot)


@tool
def generate_presentation(year: str='',keyword: str='',group_by: str='desc',snapshot_id: str='') -> str:
    """生成可编辑的销售汇报 PPT，含概览、公司/商品统计图和数据口径。

    Args:
        year: 四位年份；空为全部。
        keyword: 与查询工具相同的过滤关键词。
        group_by: 主分析分组 item/desc/date。
        snapshot_id: 复用此前办公文件的 snapshot_id；空则取最新数据。
    """
    try:
        snapshot=_snapshot(year,keyword,snapshot_id,group_by)
    except (ValueError,LookupError,FileNotFoundError) as exc:
        return '错误: '+str(exc)
    if not snapshot['rows']:
        return '（没有符合条件的已确认数据，无法生成 PPT）'
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE,XL_LABEL_POSITION
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)
    def text(slide,x,y,w,h,content,size=20,color='243A32',bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        shape.text_frame.word_wrap=True
        for i,line in enumerate(str(content).split('\n')):
            p=shape.text_frame.paragraphs[0] if i==0 else shape.text_frame.add_paragraph()
            p.text=line;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        return shape
    def page(title,dark=False):
        slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string('193C32' if dark else 'F5F6F1')
        text(slide,.7,.55,11.9,.65,title,30,'EEF2EB' if dark else '243A32',True)
        text(slide,.7,7.02,11.6,.25,'DocMind · 已确认销售数据 · '+str(len(deck.slides)),10,'B7C6BC' if dark else '727D77')
        return slide
    scope=snapshot['filters'];summary=snapshot['summary']
    slide=page('销售数据汇报',True)
    text(slide,.8,2,11.7,1.2,(scope['year'] or '全部年份')+' · '+(scope['keyword'] or '全部公司与商品'),32,'EEF2EB',True)
    text(slide,.8,4.2,11.7,1.2,'从人工核对的明细出发，形成可追溯的统计结果。\n快照时间：'+snapshot['created_at'],18,'C5D3C9')
    slide=page('数据概览')
    for i,(label,value) in enumerate([('票面金额合计',f"{summary['total_sum']:,.2f}"),('明细行数',str(summary['total_rows'])),('涉及单据',str(summary['total_docs']))]):
        text(slide,.8+i*4.1,2,3.7,.5,label,18);text(slide,.8+i*4.1,2.8,3.7,1.1,value,34,bold=True)
    text(slide,.8,5,11.7,1.2,f"数量合计：{summary['total_amount']:g}\n金额按票面字段统计，不以数量 × 单价替换；原始单据未明确币种时不标注货币单位。",18)
    names={'desc':'顾客公司','item':'商品','date':'发注日'}
    for key in dict.fromkeys([group_by,'desc','item']):
        groups=snapshot['groups'][key];shown=groups[:10]
        slide=page('按'+names[key]+'分析')
        chart_data=CategoryChartData();chart_data.categories=[g['group'][:18]+('…' if len(g['group'])>18 else '') for g in shown]
        chart_data.add_series('票面金额',[g['total'] for g in shown])
        chart=slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED,Inches(.8),Inches(1.65),Inches(8.4),Inches(4.9),chart_data).chart
        chart.has_legend=False;chart.category_axis.tick_labels.font.size=Pt(12);chart.category_axis.tick_labels.font.name='Microsoft YaHei'
        chart.value_axis.tick_labels.font.size=Pt(11)
        chart.plots[0].has_data_labels=True;chart.plots[0].data_labels.number_format='#,##0.00';chart.plots[0].data_labels.font.size=Pt(11)
        chart.plots[0].data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
        text(slide,9.5,2,3,2.7,f'共有 {len(groups)} 个分组\n展示前 {len(shown)} 组\n省略 {max(0,len(groups)-len(shown))} 组\n完整数据见 Excel。',18)
        text(slide,9.5,5,3,1,'长标签截短显示；票面负金额保留。',14)
    slide=page('数据口径与使用说明')
    text(slide,.8,1.8,11.5,3.8,'• 仅包含已人工确认的记录。\n• 年份：'+(scope['year'] or '全部')+'；关键词：'+(scope['keyword'] or '全部')+
        '\n• 数量、金额、明细行数和单据数分别统计。\n• 图表展示前10组，不代表全部分组之和。\n• 不推断因果、增长率或不存在的同期数据。\n• 同一快照可用于 Word、Excel、PPT，便于核对。',22)
    OFFICE_DIR.mkdir(parents=True,exist_ok=True);path=OFFICE_DIR/f'销售汇报_{uuid4().hex}.pptx';deck.save(path)
    return _result(path,'pptx',snapshot)
