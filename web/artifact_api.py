from pathlib import Path
from contextlib import contextmanager
from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import FileResponse,Response
from db import artifacts

router = APIRouter()


def find(ident):
    try:
        return artifacts.get(ident)
    except (LookupError,FileNotFoundError) as exc:
        raise HTTPException(404,str(exc)) from exc


@router.get('/api/artifacts/{ident}')
def info(ident: str):
    item = find(ident)
    return {key:item[key] for key in ('id','kind','name','preview_status','preview_error','created_at')}


@router.get('/api/artifacts/{ident}/download')
def download(ident: str):
    item = find(ident)
    return FileResponse(item['path'],filename=item['name'])


@router.post('/api/artifacts/{ident}/preview')
def preview(ident: str):
    find(ident)
    from services.office_preview import enqueue
    try:
        enqueue(ident)
    except RuntimeError as exc:
        raise HTTPException(429,str(exc)) from exc
    return {'ok':True}


@router.get('/api/artifacts/{ident}/preview.pdf')
def pdf(ident: str):
    item = find(ident)
    if item['preview_status']!='ready' or not item['preview_path'] or not Path(item['preview_path']).is_file():
        raise HTTPException(409,'预览尚未准备完成')
    return FileResponse(item['preview_path'],media_type='application/pdf')


@contextmanager
def open_pdf(ident):
    item=find(ident)
    if item['preview_status']!='ready' or not item['preview_path'] or not Path(item['preview_path']).is_file():
        raise HTTPException(409,'预览尚未准备完成')
    import pypdfium2
    from services.pdf_runtime import PDF_LOCK
    with PDF_LOCK:
        document=pypdfium2.PdfDocument(item['preview_path'])
        try:
            yield document
        finally:
            document.close()


@router.get('/api/artifacts/{ident}/preview/pages')
def pages(ident: str):
    with open_pdf(ident) as document:
        return {'pages':len(document)}


@router.get('/api/artifacts/{ident}/preview/page/{page_number}')
def page_image(ident: str,page_number: int,scale: float=Query(1.3,ge=.5,le=2)):
    import io
    with open_pdf(ident) as document:
        if page_number<0 or page_number>=len(document):
            raise HTTPException(404,'预览页不存在')
        page=document[page_number]
        try:
            bitmap=page.render(scale=scale)
            try:
                buffer=io.BytesIO()
                bitmap.to_pil().save(buffer,format='PNG')
                return Response(buffer.getvalue(),media_type='image/png',headers={'Cache-Control':'private, max-age=3600'})
            finally:
                bitmap.close()
        finally:
            page.close()


@router.get('/api/artifacts/{ident}/sheets')
def sheets(ident: str, sheet: int = Query(0,ge=0), offset: int = Query(0,ge=0), limit: int = Query(100,ge=1,le=200)):
    item = find(ident)
    if item['kind']!='xlsx':
        raise HTTPException(422,'此文件不是 Excel 表格')
    from openpyxl import load_workbook
    workbook = load_workbook(item['path'],read_only=True,data_only=True)
    try:
        names = workbook.sheetnames
        if sheet>=len(names):
            raise HTTPException(404,'工作表不存在')
        ws = workbook[names[sheet]]
        values = [list(row) for row in ws.iter_rows(min_row=offset+1,max_row=min(offset+limit,ws.max_row),values_only=True)] if offset<ws.max_row else []
        return {'sheets':names,'rows':values,'total_rows':ws.max_row,'offset':offset}
    finally:
        workbook.close()
